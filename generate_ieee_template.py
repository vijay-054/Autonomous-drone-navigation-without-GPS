import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START

def set_column_format(section):
    # This is a bit of a hack in python-docx to set two columns
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')[0]
    cols.set(docx.oxml.ns.qn('w:num'), '2')
    cols.set(docx.oxml.ns.qn('w:space'), '720') # 0.5 inch spacing

def create_ieee_template(filename):
    doc = docx.Document()

    # Set paper size to A4 (8.27 x 11.69 inches)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.63)
    section.right_margin = Inches(0.63)

    # Styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10)

    # Title
    title = doc.add_paragraph('Paper Title (use style: Paper Title)')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(24)
    
    # Subtitle
    subtitle = doc.add_paragraph('Subtitle as needed')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].italic = True
    
    doc.add_paragraph() # spacing

    # Authors
    authors = doc.add_paragraph('1st Given Name Surname\ndept. name of organization (of Affiliation)\nname of organization (of Affiliation)\nCity, Country\nemail address or ORCID')
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph() # spacing

    # Two columns for the rest
    section2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
    set_column_format(section2)

    # Abstract
    abs_heading = doc.add_paragraph()
    abs_run = abs_heading.add_run('Abstract—')
    abs_run.bold = True
    abs_run.italic = True
    abs_run.font.size = Pt(9)
    
    abs_text = abs_heading.add_run('This document gives formatting instructions for authors preparing papers for publication in the Proceedings of an IEEE conference. The authors must follow the instructions given in the document for the papers to be published. You can use this document as both an instruction set and as a template into which you can type your own text.')
    abs_text.bold = True
    abs_text.font.size = Pt(9)

    # Keywords
    kw_heading = doc.add_paragraph()
    kw_run = kw_heading.add_run('Keywords—')
    kw_run.bold = True
    kw_run.italic = True
    kw_run.font.size = Pt(9)
    
    kw_text = kw_heading.add_run('component, formatting, style, styling, insert')
    kw_text.font.size = Pt(9)

    doc.add_paragraph()

    # Introduction
    h1 = doc.add_paragraph('I. INTRODUCTION')
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Small caps is ideal but uppercase works for basic
    
    p = doc.add_paragraph('This document is a template. An electronic copy can be downloaded from the conference website. For questions on paper guidelines, please contact the conference publications committee as indicated on the conference website. Information about final paper submission is available from the conference website.')
    
    doc.add_paragraph('Before submitting your final paper, check that the format conforms to this template. Specifically, check the appearance of the title and author block, the appearance of section headings, document margins, column width, column spacing and other features.')
    
    h1_2 = doc.add_paragraph('II. EASE OF USE')
    h1_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    h2_1 = doc.add_paragraph('A. Maintaining the Integrity of the Specifications')
    h2_1.runs[0].italic = True
    
    doc.add_paragraph('The template is used to format your paper and style the text. All margins, column widths, line spaces, and text fonts are prescribed; please do not alter them. You may note peculiarities. For example, the head margin in this template measures proportionately more than is customary. This measurement and others are deliberate, using specifications that anticipate your paper as one part of the entire proceedings, and not as an independent document. Please do not revise any of the current designations.')

    doc.save(filename)
    print(f"Template saved to {filename}")

if __name__ == "__main__":
    create_ieee_template('IEEE_Format_Template.docx')
